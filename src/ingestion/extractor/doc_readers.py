"""Các hàm bóc tách văn bản thô từ từng định dạng file (.pdf, .docx, .xlsx, .doc, .txt)."""
from __future__ import annotations

import logging
import re
from pathlib import Path
from pypdf import PdfReader
import docx
import openpyxl

logger = logging.getLogger(__name__)


def extract_pdf(file_path: Path) -> str:
    """Bóc tách text từ file PDF sử dụng pypdf."""
    try:
        reader = PdfReader(str(file_path), strict=False)
        pages = []
        for i, page in enumerate(reader.pages):
            try:
                t = page.extract_text()
                if t and t.strip():
                    pages.append(f"--- Trang {i + 1} ---\n{t.strip()}")
            except Exception:
                continue
        return "\n\n".join(pages)
    except Exception:
        return ""


def extract_docx(file_path: Path) -> str:
    """Bóc tách text và bảng biểu từ file DOCX sử dụng python-docx."""
    try:
        doc = docx.Document(str(file_path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                unique = []
                for c in cells:
                    if not unique or c != unique[-1]:
                        unique.append(c)
                if unique:
                    parts.append(" | ".join(unique))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_xlsx(file_path: Path) -> str:
    """Bóc tách text từ các sheet bảng tính Excel (.xlsx)."""
    try:
        wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
        sheet_texts = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                valid = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if valid:
                    rows.append(" | ".join(valid))
            if rows:
                sheet_texts.append(f"=== Bảng: {sheet.title} ===\n" + "\n".join(rows))
        wb.close()
        return "\n\n".join(sheet_texts)
    except Exception:
        return ""


def extract_doc_legacy(file_path: Path) -> str:
    """Thử bóc tách text từ file .doc định dạng cũ hoặc RTF / nhị phân OLE."""
    try:
        raw_bytes = file_path.read_bytes()
        if raw_bytes.startswith(b"{\\rtf"):
            decoded = raw_bytes.decode("latin-1", errors="ignore")
            clean = re.sub(r"\\[a-z0-9]+ ?", " ", decoded)
            return re.sub(r"[{}]", "", clean)

        extracted = []
        for match in re.finditer(rb"(?:[\x20-\x7e\xa0-\xff]\x00){4,}", raw_bytes):
            try:
                chunk = match.group().decode("utf-16le", errors="ignore").strip()
                if len(chunk) >= 4 and not all(c in "-=_*~# " for c in chunk):
                    extracted.append(chunk)
            except Exception:
                pass
        return "\n".join(extracted)
    except Exception:
        return ""


def extract_txt(file_path: Path) -> str:
    """Đọc file text thô với nhiều bộ mã tiếng Việt dự phòng."""
    for enc in ["utf-8", "utf-16", "cp1258", "latin-1"]:
        try:
            return file_path.read_text(encoding=enc)
        except Exception:
            continue
    return ""
