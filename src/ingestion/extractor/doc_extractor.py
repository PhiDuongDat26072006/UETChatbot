"""Module trích xuất nội dung văn bản sạch từ các tệp tài liệu (.pdf, .docx, .xlsx, .doc, .txt)."""
from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
from contextlib import closing
from functools import lru_cache
from pathlib import Path
from typing import Any
from urllib.parse import unquote

import docx
import openpyxl
from pypdf import PdfReader
from docx.table import Table
from docx.text.paragraph import Paragraph

try:
    import pymupdf
except ImportError:  # pypdf remains a supported fallback.
    pymupdf = None

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    Presentation = None
    MSO_SHAPE_TYPE = None

from ..utils import clean_whitespace
from .date_extractor import extract_doc_date

logger = logging.getLogger(__name__)

__all__ = [
    "extract_pdf",
    "extract_docx",
    "extract_xlsx",
    "extract_pptx",
    "extract_doc_legacy",
    "extract_txt",
    "extract_document_text",
]


@lru_cache(maxsize=1)
def _ocr_language_preferences() -> tuple[str, ...]:
    """Avoid noisy missing-language errors when only English tessdata exists."""
    if shutil.which("tesseract"):
        try:
            result = subprocess.run(["tesseract", "--list-langs"], capture_output=True,
                                    text=True, timeout=5, check=True)
            installed = set(result.stdout.splitlines()[1:])
            if "vie" in installed and "eng" in installed:
                return ("vie+eng", "eng")
            if "vie" in installed:
                return ("vie",)
            if "eng" in installed:
                return ("eng",)
        except (OSError, subprocess.SubprocessError):
            pass
    return ("vie+eng", "eng")


def extract_pdf(file_path: Path) -> str:
    """Extract page text in reading order; OCR only pages without usable text."""
    if pymupdf is not None:
        try:
            with pymupdf.open(str(file_path)) as document:
                pages: list[str] = []
                for number, page in enumerate(document, start=1):
                    text = page.get_text("text", sort=True).strip()
                    if len(text) < 20:
                        for language in _ocr_language_preferences():
                            try:
                                textpage = page.get_textpage_ocr(language=language, dpi=200, full=True)
                                ocr_text = page.get_text("text", textpage=textpage, sort=True).strip()
                                if len(ocr_text) > len(text):
                                    text = ocr_text
                                break
                            except Exception as error:
                                logger.debug("OCR %s unavailable for %s page %d: %s", language, file_path, number, error)
                        if len(text) < 20:
                            logger.warning("Low text coverage in %s page %d (OCR may be needed)", file_path, number)
                    if text:
                        pages.append(f"--- Trang {number} ---\n{text}")
                    else:
                        logger.warning("No extractable text in %s page %d", file_path, number)
                return "\n\n".join(pages)
        except Exception as error:
            logger.warning("PyMuPDF failed for %s, retrying with pypdf: %s", file_path, error)

    try:
        with PdfReader(str(file_path), strict=False) as reader:
            pages: list[str] = []
            for page_number, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text()
                    if text and text.strip():
                        pages.append(f"--- Trang {page_number} ---\n{text.strip()}")
                except Exception as error:
                    logger.warning("Cannot extract %s page %s: %s", file_path, page_number, error)
                    logger.debug("PDF page failure", exc_info=True)
            return "\n\n".join(pages)
    except Exception as error:
        logger.warning("Cannot extract document %s: %s", file_path, error)
        logger.debug("Document extraction failure", exc_info=True)
        return ""


def extract_docx(file_path: Path) -> str:
    """Preserve paragraphs and tables in their original body order."""
    try:
        doc = docx.Document(str(file_path))
        parts: list[str] = []
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                if block.style and block.style.name.startswith("Heading"):
                    level = block.style.name.removeprefix("Heading").strip()
                    depth = min(int(level), 6) if level.isdigit() else 2
                    text = f"{'#' * depth} {text}"
                parts.append(text)
            elif isinstance(block, Table):
                rows = [[cell.text.strip().replace("\n", " ").replace("|", "\\|") for cell in row.cells]
                        for row in block.rows]
                if rows:
                    width = max(len(row) for row in rows)
                    rows = [row + [""] * (width - len(row)) for row in rows]
                    parts.append("\n".join(
                        ["| " + " | ".join(rows[0]) + " |",
                         "| " + " | ".join(["---"] * width) + " |"]
                        + ["| " + " | ".join(row) + " |" for row in rows[1:]]
                    ))
        return "\n\n".join(parts)
    except Exception as error:
        logger.warning("Cannot extract document %s: %s", file_path, error)
        logger.debug("Document extraction failure", exc_info=True)
        return ""


def extract_xlsx(file_path: Path) -> str:
    """Bóc tách text từ các sheet bảng tính Excel (.xlsx)."""
    try:
        with closing(openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)) as workbook:
            sheet_texts: list[str] = []
            for sheet in workbook.worksheets:
                rows: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(cell).strip().replace("|", "\\|") if cell is not None else "" for cell in row]
                    while values and not values[-1]:
                        values.pop()
                    if any(values):
                        rows.append(" | ".join(values))
                if rows:
                    sheet_texts.append(f"=== Bảng: {sheet.title} ===\n" + "\n".join(rows))
            return "\n\n".join(sheet_texts)
    except Exception as error:
        logger.warning("Cannot extract document %s: %s", file_path, error)
        logger.debug("Document extraction failure", exc_info=True)
        return ""


def extract_pptx(file_path: Path) -> str:
    """Preserve slide boundaries, textbox content, and table cell positions."""
    if Presentation is None:
        logger.warning("python-pptx is not installed: %s", file_path)
        return ""
    try:
        presentation = Presentation(str(file_path))
        slides: list[str] = []

        def render_shapes(shapes) -> list[str]:
            parts: list[str] = []
            # This is a geometric approximation, not a guaranteed reading order.
            for shape in sorted(shapes, key=lambda item: (item.top, item.left)):
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    parts.extend(render_shapes(shape.shapes))
                elif shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip().replace("\n", " ").replace("|", "\\|")
                                 for cell in row.cells]
                        parts.append(" | ".join(cells))
                elif shape.has_text_frame and shape.text.strip():
                    parts.append(shape.text.strip())
            return parts

        for number, slide in enumerate(presentation.slides, start=1):
            parts = render_shapes(slide.shapes)
            if parts:
                slides.append(f"--- Slide {number} ---\n" + "\n".join(parts))
            else:
                logger.warning("No text extracted from PPTX slide %d: %s", number, file_path)
        return "\n\n".join(slides)
    except Exception as error:
        logger.warning("Cannot extract PPTX %s: %s", file_path, error)
        return ""


def extract_doc_legacy(file_path: Path) -> str:
    """Use a real legacy DOC/RTF reader; byte regex silently breaks Vietnamese."""
    if file_path.suffix.lower() == ".rtf":
        try:
            from striprtf.striprtf import rtf_to_text
            return rtf_to_text(file_path.read_text(encoding="latin-1"))
        except (ImportError, OSError, ValueError) as error:
            logger.debug("striprtf unavailable for %s: %s", file_path, error)
    if shutil.which("soffice"):
        # Convert to DOCX in a private temp directory; never overwrite the original.
        try:
            with tempfile.TemporaryDirectory(prefix="uet-legacy-") as scratch:
                output = Path(scratch)
                profile = (output / "lo-profile").as_uri()
                subprocess.run([
                    "soffice", f"-env:UserInstallation={profile}", "--headless",
                    "--convert-to", "docx", "--outdir", str(output), str(file_path),
                ], capture_output=True, text=True, timeout=60, check=True)
                converted = output / (file_path.stem + ".docx")
                if converted.is_file():
                    return extract_docx(converted)
        except (OSError, subprocess.SubprocessError) as error:
            logger.warning("LibreOffice conversion failed for %s: %s", file_path, error)
    if not shutil.which("antiword"):
        logger.warning("Install LibreOffice or antiword to extract legacy document: %s", file_path)
        return ""
    try:
        result = subprocess.run(["antiword", "-m", "UTF-8.txt", str(file_path)],
                                capture_output=True, text=True, timeout=30, check=True)
        return result.stdout
    except (OSError, subprocess.SubprocessError) as error:
        logger.warning("Cannot extract legacy DOC %s: %s", file_path, error)
        return ""


def extract_txt(file_path: Path) -> str:
    """Đọc file text thô với nhiều bộ mã tiếng Việt dự phòng."""
    for enc in ["utf-8", "utf-16", "cp1258", "latin-1"]:
        try:
            return file_path.read_text(encoding=enc)
        except UnicodeError:
            logger.debug("Cannot decode %s with %s", file_path, enc)
        except OSError as error:
            logger.warning("Cannot read %s: %s", file_path, error)
            return ""
    return ""


def extract_document_text(file_path: Path | str, min_length: int = 30) -> dict[str, Any] | None:
    """
    Nhận diện định dạng và trích xuất nội dung văn bản sạch từ một file tài liệu.

    Hỗ trợ PDF, DOCX, XLSX, DOC, TXT/MD/CSV. Trích xuất cả ngày ban hành từ tên file
    hoặc nội dung.
    """
    path = Path(file_path)
    if not path.is_file() or path.name.startswith(".") or path.name in ("files_list.txt", "endpoints.txt"):
        return None

    ext = path.suffix.lower()
    raw_text = ""

    if ext == ".pdf":
        raw_text = extract_pdf(path)
    elif ext == ".docx":
        raw_text = extract_docx(path)
    elif ext in (".xlsx", ".xlsm"):
        raw_text = extract_xlsx(path)
    elif ext == ".pptx":
        raw_text = extract_pptx(path)
    elif ext in (".doc", ".rtf"):
        raw_text = extract_doc_legacy(path)
    elif ext in (".txt", ".md", ".csv"):
        raw_text = extract_txt(path)
    else:
        return None

    cleaned_text = clean_whitespace(raw_text)
    if len(cleaned_text) < min_length:
        return None

    readable_filename = unquote(path.name)
    published_date = extract_doc_date(file_path=path, text=cleaned_text)

    return {
        "filename": readable_filename,
        "file_path": str(path),
        "file_type": ext.lstrip("."),
        "published_date": published_date,
        "text": cleaned_text,
        "text_length": len(cleaned_text),
    }
